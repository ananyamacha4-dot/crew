"""Convert CREWAI_GUIDE.md into a shareable .docx with proper formatting."""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = Path(__file__).parent / "CREWAI_GUIDE.md"
DST = Path(__file__).parent / "CrewAI_Developer_Guide.docx"


# ---------- inline formatting (bold / italic / code / link) ----------

INLINE_RE = re.compile(
    r"(\*\*([^*]+)\*\*)"        # **bold**
    r"|(\*([^*]+)\*)"            # *italic*
    r"|(`([^`]+)`)"              # `code`
    r"|(\[([^\]]+)\]\(([^)]+)\))"  # [text](url)
)


def add_runs(paragraph, text):
    """Parse inline markdown and add formatted runs to a paragraph."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        if m.group(2):  # bold
            r = paragraph.add_run(m.group(2))
            r.bold = True
        elif m.group(4):  # italic
            r = paragraph.add_run(m.group(4))
            r.italic = True
        elif m.group(6):  # inline code
            r = paragraph.add_run(m.group(6))
            r.font.name = "Consolas"
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        elif m.group(7):  # link
            add_hyperlink(paragraph, m.group(8), m.group(9))
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_hyperlink(paragraph, text, url):
    """Add a clickable hyperlink to a paragraph."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)

    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


# ---------- block-level builders ----------

def add_code_block(doc, lines, lang=""):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    shade(p, "F2F2F2")
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)


def shade(paragraph, hex_fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pPr.append(shd)


def add_table(doc, rows):
    if not rows:
        return
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Light Grid Accent 1"
    for r_idx, row in enumerate(rows):
        for c_idx, cell_text in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            add_runs(p, cell_text.strip())
            if r_idx == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()


def add_heading(doc, text, level):
    h = doc.add_heading(level=min(level, 4))
    add_runs(h, text)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    add_runs(p, text)


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    add_runs(p, text)


def add_quote(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.4)
    shade(p, "FFF8E1")
    run = p.add_run(text)
    run.italic = True


def add_hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "999999")
    pBdr.append(bottom)
    pPr.append(pBdr)


# ---------- main parser ----------

def convert(src_path: Path, dst_path: Path):
    text = src_path.read_text(encoding="utf-8")
    lines = text.split("\n")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i]

        # fenced code block
        if line.startswith("```"):
            lang = line[3:].strip()
            i += 1
            buf = []
            while i < len(lines) and not lines[i].startswith("```"):
                buf.append(lines[i])
                i += 1
            add_code_block(doc, buf, lang)
            i += 1
            continue

        # table (line with | and next line is separator)
        if "|" in line and i + 1 < len(lines) and re.match(r"^\s*\|?[\s\-:|]+\|?\s*$", lines[i + 1]) and "-" in lines[i + 1]:
            rows = []
            header = [c.strip() for c in line.strip().strip("|").split("|")]
            rows.append(header)
            i += 2  # skip separator
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(row)
                i += 1
            add_table(doc, rows)
            continue

        # horizontal rule
        if re.match(r"^---+\s*$", line):
            add_hr(doc)
            i += 1
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            add_heading(doc, m.group(2).strip(), len(m.group(1)))
            i += 1
            continue

        # blockquote
        if line.startswith(">"):
            add_quote(doc, line.lstrip("> ").strip())
            i += 1
            continue

        # bullet list (- or *)
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            indent = len(m.group(1)) // 2
            add_bullet(doc, m.group(2), level=indent)
            i += 1
            continue

        # numbered list
        m = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if m:
            add_numbered(doc, m.group(1))
            i += 1
            continue

        # blank line
        if not line.strip():
            i += 1
            continue

        # plain paragraph (gather wrapped lines)
        para_lines = [line]
        j = i + 1
        while j < len(lines) and lines[j].strip() and not re.match(
            r"^(#|>|```|---|\s*[-*]\s|\s*\d+\.\s|\|)", lines[j]
        ):
            para_lines.append(lines[j])
            j += 1
        p = doc.add_paragraph()
        add_runs(p, " ".join(l.strip() for l in para_lines))
        i = j

    doc.save(dst_path)
    print(f"Wrote: {dst_path}")


if __name__ == "__main__":
    convert(SRC, DST)
